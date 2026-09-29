"""Extrai texto puro de cada tipo de arquivo suportado."""

from html.parser import HTMLParser
from pathlib import Path


class _HTMLText(HTMLParser):
    SKIP = {"script", "style", "noscript"}
    BLOCK = {"p", "div", "br", "li", "tr", "section", "article"}

    def __init__(self):
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip += 1
        elif tag in {"h1", "h2", "h3", "h4"}:
            self.parts.append("\n" + "#" * int(tag[1]) + " ")
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self._skip:
            self._skip -= 1
        elif tag in {"h1", "h2", "h3", "h4"}:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def _pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        try:
            t = page.extract_text() or ""
        except Exception:
            t = ""
        if t.strip():
            pages.append(f"## Página {i}\n\n{t.strip()}")
    return "\n\n".join(pages)


def _docx(path: Path) -> str:
    import docx

    doc = docx.Document(str(path))
    out = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading") or style.startswith("título"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            out.append("#" * min(int(level), 4) + " " + text)
        else:
            out.append(text)
    for table in doc.tables:
        for row in table.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n\n".join(out)


def _html(path: Path) -> str:
    p = _HTMLText()
    p.feed(path.read_text(encoding="utf-8", errors="replace"))
    return "".join(p.parts)


def extract_text(path: Path, raw: bytes) -> str | None:
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            return _pdf(path)
        if ext == ".docx":
            return _docx(path)
        if ext in (".html", ".htm"):
            return _html(path)
    except Exception:
        return None
    if b"\x00" in raw[:8192]:
        return None  # binário disfarçado
    return raw.decode("utf-8", errors="replace")
